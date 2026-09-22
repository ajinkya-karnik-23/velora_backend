"""Build the completed TWP work paper for a control from its client data.

Takes the empty work-paper template and fills it from the control definition,
the sampling determination and the per-sample testing output — producing the
workbook the Controls view serves once every sample has been tested.

Re-run after editing any of the source JSONs:

    poetry run python -m scripts.build_twp_completed IA8.CA02 19A1

Values written into dropdown-validated cells are taken verbatim from the
template's own lookup lists (rows 176-194, the "do not touch these values"
block), so Excel does not flag them as invalid.
"""

from __future__ import annotations

import sys
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.worksheet.worksheet import Worksheet

from app.core.config import settings
from app.services.control_matching import (
    find_control_json_by_control_and_entity,
    load_control_json,
)
from app.services.evidence_verification import (
    build_methodology,
    expected_filename_for_sample,
    load_filename_rule,
)
from app.services.sampling_matrix import (
    calculate_sample_size,
    extract_control_parameters,
    load_sampling_matrix,
    load_sampling_notes,
)
from app.services.test_output_matching import build_samples, find_test_output
from app.services.twp_sample_details import add_sample_details_sheet
from app.services.twp_templates import load_twp_templates
from app.utils.workbook_annotations import flag_out_of_scope_fields

# The template's own frequency list (B188:B194) — a control's raw frequency is
# mapped onto one of these so the cell stays valid.
FREQUENCY_CHOICES = [
    "Annual",
    "Bi-annual",
    "Quarterly",
    "Monthly",
    "Weekly",
    "Daily",
    "Multiple times per day",
]

RESULT_OK = "OK"
RESULT_NOT_OK = "Not OK"
RESULT_NA = "N/A"
# A sample whose assessment is not finished. Written into the template's own
# result list (the empty E187 slot of E184:E187) so the dropdown accepts it.
RESULT_PENDING = "Pending"
RESULT_LIST_SLOT = "E187"

TEST_RESULT_EFFECTIVE = "EFFECTIVE - Test executed without exception(s)"
TEST_RESULT_INEFFECTIVE = "INEFFECTIVE - Test executed with exception(s)"
# Already in the template's Test result list (I179).
TEST_RESULT_NOT_TESTED = "NOT TESTED - Test result not available"

# The per-sample results band. Every row here carries a validated result cell
# (E118:E143 plus E144), giving exactly 27 rows.
RESULTS_FIRST_ROW = 118
RESULTS_LAST_ROW = 144


def _set(ws: Worksheet, coord: str, value: Any) -> None:
    """Write a value, redirecting to the anchor when the cell is merged."""
    for rng in ws.merged_cells.ranges:
        if coord in rng:
            ws.cell(row=rng.min_row, column=rng.min_col).value = value
            return
    ws[coord] = value


def _map_frequency(raw: str | None) -> str:
    """Fit a control's frequency onto the template's allowed list."""
    text = (raw or "").strip().lower()
    for choice in FREQUENCY_CHOICES:
        if choice.lower() == text:
            return choice
    # "Upon Occurrence" and similar have no slot in the list; the closest
    # honest reading of an event-driven control is its highest cadence.
    return "Multiple times per day" if text else ""


# Parameters that best identify a sample. Controls describing their samples
# differently fall back to whatever they do carry.
PREFERRED_LABELS = (
    "Posting date",
    "Amount in local cur.",
    "Type",
    # Payment controls name the same facts slightly differently.
    "Document Number",
    "Posting Date",
    "Amount in local currency",
)

# Long narrative fields are never useful as a row label.
LABEL_MAX_LEN = 60


def _sample_label(sample: dict) -> str:
    """One results-row description: sample number, document, key parameters."""
    parts = [f"Sample {sample.get('sample_no')}"]
    if sample.get("document_no"):
        parts.append(f"Doc {sample['document_no']}")

    params = sample.get("parameters") or []
    chosen = [p for p in params if p.get("label") in PREFERRED_LABELS]
    if not chosen:
        # This control names its samples some other way (a customer, a month).
        # Take the first few short parameters so the row still identifies
        # itself instead of reading "Sample 1" with nothing else.
        chosen = [
            p
            for p in params
            if p.get("value") not in (None, "") and len(str(p.get("value"))) <= LABEL_MAX_LEN
        ][:3]

    for param in chosen:
        value = param.get("value")
        if value not in (None, ""):
            parts.append(f"{param['label']}: {value}")
    return " | ".join(parts)


def _sample_fields(sample: dict) -> dict:
    """Flatten a normalised sample back to field -> value, as the filename
    rule is written against the client's own field names."""
    fields = {k: v for k, v in sample.items() if k != "parameters"}
    for param in sample.get("parameters") or []:
        label = param.get("label")
        if label:
            fields[label] = param.get("value")
    return fields


def _population_split(methodology: dict) -> str:
    """How the tested population divides, phrased for the conclusion line.

    Controls using the journal-entry methodology keep their original wording;
    any control declaring its own categories is described by those instead.
    """
    if methodology.get("manual_entries") is not None:
        return f"{methodology['manual_entries']} manual, {methodology['nr_entries']} NR"
    lines = methodology.get("breakdown") or []
    return ", ".join(f"{line['value']} {line['label'].lower()}" for line in lines)


def _evidence_by_sample(rule: dict | None, samples: list[dict]) -> list[list[str]]:
    """The files each sample was tested against.

    A sample's own `evidences_used` wins, since it names what the test actually
    read. Otherwise the control's filename rule says what it had to be run on:
    one file for every sample, or one derived from the sample's own fields.
    """
    fixed = expected_filename_for_sample(rule, None) if rule else None
    result: list[list[str]] = []
    for sample in samples:
        cited = [str(n).strip() for n in sample.get("evidences_used") or [] if str(n).strip()]
        if not cited and rule:
            name = fixed or expected_filename_for_sample(rule, _sample_fields(sample))
            cited = [name] if name else []
        result.append(cited)
    return result


def _is_pending(sample: dict) -> bool:
    return (sample.get("result") or "").strip().upper() == "PENDING"


# The summary box is D147:I148, about 110 characters of 12pt Calibri per line.
SUMMARY_CHARS_PER_LINE = 105
SUMMARY_LINE_POINTS = 16.0


def _fit_summary_box(ws: Worksheet, first_row: int, text: str) -> None:
    """Grow the summary box's lower row so the whole narrative shows without
    the reviewer resizing it. Never shrinks below the template's own height."""
    lines = sum(max(1, -(-len(part) // SUMMARY_CHARS_PER_LINE)) for part in text.split("\n"))
    needed = lines * SUMMARY_LINE_POINTS + 6
    top = ws.row_dimensions[first_row].height or 15.65
    bottom = ws.row_dimensions[first_row + 1].height or 15.65
    if top + bottom < needed:
        ws.row_dimensions[first_row + 1].height = needed - top


def _root_cause(exceptions: list[dict], payload: dict | None) -> str:
    """Why samples did not pass, in terms the control's own data supports.

    The journal-entry control keeps its established wording. A control that
    declares its own methodology is described from the exceptions themselves:
    which samples, and the result each returned, rather than a reason the
    testing output never gave.
    """
    declared = ((payload or {}).get("control_test_output") or {}).get("methodology")
    if not isinstance(declared, dict):
        return (
            "Journal entries flagged NR could not be validated against the "
            "supporting documentation provided."
        )
    by_result: dict[str, list[str]] = {}
    for sample in exceptions:
        result = str(sample.get("result") or "no result").strip().upper().replace("_", " ")
        by_result.setdefault(result, []).append(str(sample.get("sample_no")))
    parts = [
        f"{len(nos)} sample(s) returned {result.lower()} (samples {', '.join(nos)})"
        for result, nos in by_result.items()
    ]
    return "; ".join(parts) + ". See the Sample Testing Details sheet for each sample's validation."


def build(control_number: str, entity_code: str) -> Path:
    templates = load_twp_templates(
        Path(settings.TWP_TEMPLATE_MAP_PATH), control_number, entity_code
    )
    source = Path(templates.empty or settings.TWP_TEMPLATE_PATH)
    # Unmapped controls land beside the template they were built from, inside
    # whichever client-data folder is configured.
    target = Path(
        templates.completed
        or Path(settings.TWP_TEMPLATE_PATH).parent
        / f"{control_number}.{entity_code}.TWP.Completed.xlsm"
    )
    if not source.exists():
        raise SystemExit(f"Empty template not found: {source}")

    control_path = find_control_json_by_control_and_entity(
        Path(settings.CONTROL_JSONS_PATH), control_number, entity_code
    )
    if control_path is None:
        raise SystemExit(f"No control JSON for {control_number}.{entity_code}")
    definition = load_control_json(control_path)
    rcm = definition.get("rcm_details") or {}
    # Audit narrative for the "Summary of test performed" box: nature of the
    # testing, then results and root cause. Written per control in its JSON.
    template_summary = str(
        (definition.get("control_details") or {}).get("template_summary") or ""
    ).strip()

    payload = find_test_output(Path(settings.TEST_OUTPUTS_PATH), control_number, entity_code)
    samples = build_samples(payload) if payload else []
    methodology = build_methodology(samples, (payload or {}).get("control_test_output"))

    params = extract_control_parameters({"rcm_details": rcm})
    matrix = load_sampling_matrix(Path(settings.SAMPLING_MATRIX_PATH))
    sample_size = calculate_sample_size(
        matrix,
        risk_level=params["risk_level"],
        frequency=params["frequency"],
        phase=params["phase"],
    ) or str(len(samples))
    notes = load_sampling_notes(Path(settings.SAMPLING_METADATA_PATH), control_number, entity_code)

    # Pending samples are not yet concluded, so they are neither passes nor
    # exceptions; they hold the overall result open instead.
    pending = [s for s in samples if _is_pending(s)]
    exceptions = [
        s
        for s in samples
        if (s.get("result") or "").upper() not in ("PASS", "PASSED") and not _is_pending(s)
    ]

    wb = openpyxl.load_workbook(source, keep_vba=True)
    ws = wb["Template"]

    # ── Work paper preparation ────────────────────────────────────────────
    _set(ws, "D5", rcm.get("Entity Name") or entity_code)
    _set(ws, "D7", control_number)
    _set(ws, "D9", rcm.get("Title"))
    _set(ws, "D11", f"TWP-{control_number}.{entity_code}")
    _set(ws, "D13", rcm.get("Description"))
    _set(ws, "D15", "Yes" if str(rcm.get("IPE Name of Report")).strip().lower() == "yes" else "No")
    _set(ws, "D17", "Validated through Usage")

    # ── Test plan ─────────────────────────────────────────────────────────
    _set(ws, "D22", len(samples) or sample_size)
    _set(ws, "F22", (notes.get("frequency") or {}).get("text") if notes.get("frequency") else None)
    _set(ws, "H24", (params["risk_level"] or "").upper() or None)
    _set(ws, "D28", _map_frequency(params["frequency"]))
    _set(ws, "D30", rcm.get("Effective Date"))
    _set(ws, "D32", rcm.get("Control Type") or "Manual")
    _set(ws, "D34", rcm.get("Control Objective"))

    # ── Test execution ────────────────────────────────────────────────────
    _set(ws, "D114", rcm.get("Control Owner"))

    # Rows are never inserted here: openpyxl shifts cell values but leaves
    # merged ranges where they are, which would silently desynchronise every
    # label below the band from its value. Overflow is refused instead, so the
    # template gets extended deliberately rather than producing a broken
    # workbook.
    capacity = RESULTS_LAST_ROW - RESULTS_FIRST_ROW + 1
    if len(samples) > capacity:
        raise SystemExit(
            f"{len(samples)} samples exceed the template's {capacity}-row results "
            f"band (rows {RESULTS_FIRST_ROW}-{RESULTS_LAST_ROW}). Extend the band "
            "in the template — copying the result-cell validation down — then "
            "update RESULTS_LAST_ROW."
        )

    for offset, sample in enumerate(samples):
        row = RESULTS_FIRST_ROW + offset
        passed = (sample.get("result") or "").upper() in ("PASS", "PASSED")
        _set(ws, f"D{row}", _sample_label(sample))
        _set(
            ws,
            f"E{row}",
            RESULT_OK if passed else RESULT_PENDING if _is_pending(sample) else RESULT_NOT_OK,
        )
        _set(ws, f"F{row}", None if passed else sample.get("validation"))

    if pending:
        ws[RESULT_LIST_SLOT] = RESULT_PENDING
        # E144 alone validates against E183:E186, one row short of the slot.
        for dv in ws.data_validations.dataValidation:
            if "E144" in str(dv.sqref) and dv.formula1 == "$E$183:$E$186":
                dv.formula1 = "$E$183:$E$187"

    # Conclusion block. Fixed positions: the band is never resized, so these
    # sit exactly where the template puts them. The summary caption lives in
    # column D with its value beneath; the rest are captioned in column B with
    # the value in D on the same row.
    summary_label = 146
    headline = (
        f"{methodology['methodology']}: {methodology['total_samples']} samples tested "
        f"({_population_split(methodology)}). "
        f"{len(samples) - len(exceptions) - len(pending)} passed, "
        f"{len(exceptions)} exception(s)" + (f", {len(pending)} pending." if pending else ".")
    )
    summary_text = f"{headline}\n\n{template_summary}" if template_summary else headline
    _set(ws, f"D{summary_label + 1}", summary_text)
    _fit_summary_box(ws, summary_label + 1, summary_text)
    _set(ws, f"D{summary_label + 4}", len(exceptions))
    if exceptions:
        _set(ws, f"D{summary_label + 6}", _root_cause(exceptions, payload))
    # An exception decides the result on its own; otherwise pending samples
    # leave it open rather than letting it read as effective.
    _set(
        ws,
        f"D{summary_label + 8}",
        RESULT_NOT_OK if exceptions else RESULT_PENDING if pending else RESULT_OK,
    )
    _set(
        ws,
        f"D{summary_label + 10}",
        (
            TEST_RESULT_INEFFECTIVE
            if exceptions
            else TEST_RESULT_NOT_TESTED if pending else TEST_RESULT_EFFECTIVE
        ),
    )

    # Evidence collected — the work paper records what the test was run against.
    # A control either names one file for every sample, one per sample derived
    # from that sample's own data, or — when it is not filename-gated at all —
    # the files each sample's testing output says it was validated on.
    rule = load_filename_rule(
        Path(settings.EVIDENCE_FILENAME_MAP_PATH), control_number, entity_code
    )
    fixed = expected_filename_for_sample(rule, None) if rule else None
    if fixed:
        _set(ws, f"D{summary_label + 13}", fixed)
        _set(
            ws,
            f"E{summary_label + 13}",
            f"Test of effectiveness evidence supporting all {len(samples)} samples.",
        )
    elif rule:
        per_sample = [
            expected_filename_for_sample(rule, _sample_fields(sample)) for sample in samples
        ]
        named = [name for name in per_sample if name]
        if named:
            _set(ws, f"D{summary_label + 13}", ", ".join(named))
            _set(
                ws,
                f"E{summary_label + 13}",
                f"One evidence file per sample, {len(named)} in total.",
            )
    else:
        # Ungated control: the testing output itself lists what each sample was
        # validated on. Names are de-duplicated in first-seen order, since one
        # file can support several samples.
        cited: list[str] = []
        for sample in samples:
            for name in sample.get("evidences_used") or []:
                text = str(name).strip()
                if text and text not in cited:
                    cited.append(text)
        if cited:
            _set(ws, f"D{summary_label + 13}", ", ".join(cited))
            _set(
                ws,
                f"E{summary_label + 13}",
                f"Evidence cited across {len(samples)} samples, {len(cited)} file(s) in total.",
            )

    # Fields the automation does not answer are handed to the reviewer rather
    # than left reading "out-of-scope".
    flag_out_of_scope_fields(ws)

    # The full per-sample record behind the results band above.
    add_sample_details_sheet(
        wb,
        samples,
        _evidence_by_sample(rule, samples),
        control_number=control_number,
        entity_code=entity_code,
    )

    target.parent.mkdir(parents=True, exist_ok=True)
    wb.save(target)
    return target


def main() -> None:
    control_number = sys.argv[1] if len(sys.argv) > 1 else "IA8.CA02"
    entity_code = sys.argv[2] if len(sys.argv) > 2 else "19A1"
    out = build(control_number, entity_code)
    print(f"wrote {out} ({out.stat().st_size // 1024} KB)")


if __name__ == "__main__":
    main()
