"""Worksheet annotations applied to generated work papers.

The TWP template carries fields the automation deliberately does not answer —
the IPE questionnaire, for one, which the reviewer fills in from their own
inquiries. Left as they ship, those cells read "out-of-scope", which says
nothing about who is expected to act on them.

`flag_out_of_scope_fields` turns each one into a highlighted, italic prompt
so the reviewer can see at a glance what is still theirs to complete.
"""

from __future__ import annotations

from typing import TYPE_CHECKING

from openpyxl.styles import Font, PatternFill

if TYPE_CHECKING:
    from openpyxl.worksheet.worksheet import Worksheet

# Cell text that marks a field as outside the automation's scope, compared
# case- and whitespace-insensitively.
OUT_OF_SCOPE_MARKERS = ("out-of-scope", "out of scope", "outofscope")

# What the reviewer sees in place of the marker.
COMMENT_PROMPT = "Please add your comments"

# Excel's own "Neutral" pairing — a highlighter yellow that stays readable in
# print and does not collide with the result cells' own conditional colours.
HIGHLIGHT_FILL = "FFEB9C"
HIGHLIGHT_TEXT = "9C6500"


def _is_flaggable(value: object, prompt: str, markers: tuple[str, ...]) -> bool:
    """Whether this cell is an out-of-scope field, or already flagged as one.

    Re-flagging an already-flagged cell is a no-op rather than an error, so
    the annotation can run on a workbook built from an annotated template.
    """
    if not isinstance(value, str):
        return False
    text = " ".join(value.split()).strip().lower()
    return text == prompt.lower() or text in markers


def flag_out_of_scope_fields(
    worksheet: Worksheet,
    prompt: str = COMMENT_PROMPT,
    markers: tuple[str, ...] = OUT_OF_SCOPE_MARKERS,
) -> list[str]:
    """Highlight every out-of-scope field and prompt the reviewer to fill it.

    Each matching cell is given the prompt in italics on a yellow ground.
    Only the anchor of a merged field is written and styled: it is the one
    cell in the range that can hold a value, and Excel paints a merged range
    from its anchor, so the whole band picks up the highlight.

    Returns the coordinates flagged, in row order, so callers can report what
    was marked (and notice when a template stops declaring any).
    """
    fill = PatternFill(fill_type="solid", start_color=HIGHLIGHT_FILL, end_color=HIGHLIGHT_FILL)

    flagged: list[str] = []
    for row in worksheet.iter_rows():
        for cell in row:
            if not _is_flaggable(cell.value, prompt, markers):
                continue
            cell.value = prompt
            cell.fill = fill
            # Keep the cell's own face and size; only the italics and colour
            # are the annotation's to decide.
            cell.font = Font(
                name=cell.font.name,
                size=cell.font.size,
                bold=cell.font.bold,
                italic=True,
                color=HIGHLIGHT_TEXT,
            )
            flagged.append(cell.coordinate)
    return flagged
