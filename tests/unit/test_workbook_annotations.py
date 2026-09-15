"""Out-of-scope field annotation on generated work papers."""

from __future__ import annotations

import openpyxl

from app.utils.workbook_annotations import (
    COMMENT_PROMPT,
    HIGHLIGHT_FILL,
    flag_out_of_scope_fields,
)


def _sheet() -> openpyxl.worksheet.worksheet.Worksheet:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["B2"] = "Reporting entity:"
    ws["D2"] = "19A1 ArcelorMittal Exports DMCC"
    ws["B3"] = "What is the source data of the report?"
    ws["D3"] = "out-of-scope"
    ws["D4"] = "Out of Scope"
    return ws


def test_flags_every_out_of_scope_field() -> None:
    ws = _sheet()

    assert flag_out_of_scope_fields(ws) == ["D3", "D4"]

    for coord in ("D3", "D4"):
        assert ws[coord].value == COMMENT_PROMPT
        assert ws[coord].font.italic is True
        assert ws[coord].fill.start_color.rgb.endswith(HIGHLIGHT_FILL)


def test_leaves_answered_fields_alone() -> None:
    ws = _sheet()
    flag_out_of_scope_fields(ws)

    assert ws["D2"].value == "19A1 ArcelorMittal Exports DMCC"
    assert ws["D2"].font.italic in (None, False)
    assert ws["B3"].value == "What is the source data of the report?"


def test_reflagging_is_a_no_op() -> None:
    """A workbook built from an annotated template flags the same cells."""
    ws = _sheet()
    flag_out_of_scope_fields(ws)

    assert flag_out_of_scope_fields(ws) == ["D3", "D4"]
    assert ws["D3"].value == COMMENT_PROMPT
